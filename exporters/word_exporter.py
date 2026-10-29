"""Create a detailed, plain-language Microsoft Word report for scan results."""
from datetime import datetime
from typing import List

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor

from scanner.models import Host


def _severity(cves: list) -> str:
    levels = {str(c.get("severity", "")).upper() for c in cves if isinstance(c, dict)}
    for level in ("CRITICAL", "HIGH", "MEDIUM", "LOW"):
        if level in levels:
            return level
    return "No CVE hint"


def _plain_state(state: str) -> str:
    meanings = {
        "OPEN": "Open — a service responded on this port.",
        "CLOSED": "Closed — the device responded, but no service accepted the connection.",
        "FILTERED": "Filtered — a firewall or network rule may be preventing a clear answer.",
        "OPEN|FILTERED": "Open or filtered — the scan could not reliably distinguish the two.",
    }
    return meanings.get(state.upper(), f"{state} — see the technical result for details.")


def export_word_report(
    hosts: List[Host],
    filepath: str,
    target: str = "Local Subnet",
    scan_time: float = 0.0,
    scan_mode: str = "Not recorded",
    port_selection: str = "Not recorded",
) -> str:
    """Write a detailed, non-technical .docx report and return its path."""
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.75)
    section.right_margin = Inches(0.75)

    styles = doc.styles
    styles["Normal"].font.name = "Aptos"
    styles["Normal"].font.size = Pt(10)
    styles["Title"].font.name = "Aptos Display"
    styles["Title"].font.size = Pt(26)
    styles["Title"].font.color.rgb = RGBColor(31, 78, 121)
    for style_name in ("Heading 1", "Heading 2"):
        styles[style_name].font.name = "Aptos Display"
        styles[style_name].font.color.rgb = RGBColor(31, 78, 121)

    total_open = sum(len(host.open_ports) for host in hosts)
    all_cves = [c for h in hosts for p in h.ports for c in p.cve_list if isinstance(c, dict)]
    severity_counts = {level: sum(1 for c in all_cves if str(c.get("severity", "")).upper() == level)
                       for level in ("CRITICAL", "HIGH", "MEDIUM", "LOW")}

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("LAN SCANNER\nNETWORK CHECK REPORT")
    run.bold = True
    run.font.size = Pt(23)
    run.font.color.rgb = RGBColor(31, 78, 121)

    p = doc.add_paragraph("Plain-language summary and technical findings")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.runs[0].italic = True

    doc.add_paragraph(f"Report generated: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}")
    doc.add_paragraph(f"Target checked: {target}")
    doc.add_paragraph(f"Scan method: {scan_mode}")
    doc.add_paragraph(f"Port selection: {port_selection}")
    doc.add_paragraph(f"Scan duration: {scan_time:.2f} seconds")

    doc.add_heading("1. Summary for non-technical readers", level=1)
    summary = doc.add_table(rows=0, cols=2)
    summary.style = "Light Shading Accent 1"
    for label, value in [
        ("Devices found", str(len(hosts))),
        ("Open ports found", str(total_open)),
        ("Potential vulnerability hints", str(len(all_cves))),
        ("Critical severity hints", str(severity_counts["CRITICAL"])),
        ("High severity hints", str(severity_counts["HIGH"])),
    ]:
        cells = summary.add_row().cells
        cells[0].text = label
        cells[1].text = value

    if not hosts:
        doc.add_paragraph("No devices were returned by the scan. This may mean the target was unreachable, discovery was blocked, or no hosts responded; it does not prove the network is secure.")
    elif total_open == 0 and not all_cves:
        doc.add_paragraph("The scan did not identify open ports or vulnerability hints in the results collected. This is not a guarantee that the device or network is secure: the scan only checks the selected targets, ports, and methods.")
    else:
        doc.add_paragraph(f"The scan found {len(hosts)} device(s) and {total_open} open port(s). An open port means a service appears to be listening; it is not automatically a security problem. Review the findings below and confirm that each service is expected and appropriately protected.")

    doc.add_heading("2. What the results mean", level=1)
    definitions = [
        ("Device / host", "A computer, phone, server, printer, or other network-connected device that responded to discovery or scanning."),
        ("Port", "A numbered connection point used by network services. For example, a web server may use port 80 or 443."),
        ("Open", "A service responded on that port. Check whether the service is needed and who can reach it."),
        ("Closed", "The device responded, but no application accepted the connection on that port at scan time."),
        ("Filtered", "A firewall or network rule may have blocked the probe, so the scanner could not determine the port state clearly."),
        ("CVE / vulnerability hint", "A possible match based on service information or a banner. It is not proof that the device is vulnerable; verify versions and vendor advisories."),
        ("No findings", "Means this scan did not detect a matching issue. It does not mean the device has no vulnerabilities."),
    ]
    for term, definition in definitions:
        p = doc.add_paragraph(style="List Bullet")
        r = p.add_run(f"{term}: ")
        r.bold = True
        p.add_run(definition)

    doc.add_heading("3. Device-by-device results", level=1)
    if not hosts:
        doc.add_paragraph("No host details are available for this scan.")
    for host in hosts:
        doc.add_heading(f"Device: {host.ip}", level=2)
        doc.add_paragraph(f"Host name: {host.hostname or 'Not identified'} | MAC address: {host.mac or 'Not identified'} | Vendor: {host.vendor or 'Unknown'}")
        doc.add_paragraph(f"Open ports on this device: {len(host.open_ports)}")
        if not host.ports:
            doc.add_paragraph("No port checks were recorded for this device. This may occur in host-discovery-only mode.")
            continue
        table = doc.add_table(rows=1, cols=5)
        table.style = "Light Shading Accent 1"
        headers = ["Port / protocol", "State", "Service", "Possible CVE hints", "Plain-language meaning"]
        for i, heading in enumerate(headers):
            table.rows[0].cells[i].text = heading
        for port in host.ports:
            cves = [c for c in port.cve_list if isinstance(c, dict)]
            cve_text = ", ".join(str(c.get("cve_id", "Unknown ID")) + " (" + str(c.get("severity", "unspecified")) + ")" for c in cves) or "None detected"
            meaning = _plain_state(port.state)
            if port.state.upper() == "OPEN":
                meaning += " Confirm this service is needed and access is restricted appropriately."
            cells = table.add_row().cells
            values = [f"{port.port}/{port.protocol.upper()}", port.state, port.service or "Unknown", cve_text, meaning]
            for i, value in enumerate(values):
                cells[i].text = str(value)
            if cves:
                p = doc.add_paragraph()
                p.add_run(f"Details for {port.port}/{port.protocol.upper()}: ").bold = True
                for cve in cves:
                    p.add_run(f"{cve.get('cve_id', 'CVE hint')} — {cve.get('title', cve.get('description', 'No description available'))}. ")
                    fix = cve.get("fix")
                    if fix:
                        p.add_run(f"Suggested next step: {fix} ")

    doc.add_heading("4. Recommended next steps", level=1)
    recommendations = []
    if total_open:
        recommendations.append("Review each open port with the device owner. Disable services that are not needed and limit access to required users or networks.")
    if all_cves:
        recommendations.append("Treat CVE entries as leads to investigate, not confirmed vulnerabilities. Verify the exact product and version against the vendor's security advisory or an approved vulnerability assessment.")
    recommendations.extend([
        "Keep operating systems and network-facing applications updated with supported security patches.",
        "Use firewalls and access controls to limit who can connect to administrative or database services.",
        "Repeat the scan after changes and compare reports to check whether the exposure has changed.",
    ])
    for item in dict.fromkeys(recommendations):
        doc.add_paragraph(item, style="List Number")

    doc.add_heading("5. Scope and limitations", level=1)
    doc.add_paragraph(f"This report covers target scope '{target}' using the scan configuration recorded above. Results depend on the selected scan technique, ports, timeouts, network reachability, and permissions. A port scan is not a full security audit or penetration test. Banner-based CVE matching can produce false positives or miss issues. No result should be treated as a guarantee of security.")
    doc.add_paragraph("Authorized use only: scan systems and networks that you own or have explicit permission to assess.")

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run("LAN Scanner & Active Port Scanner | Educational and authorized testing use only")

    doc.save(filepath)
    return filepath
